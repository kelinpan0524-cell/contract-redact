#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""脱敏自检：把已经脱敏的文件再跑一遍同一套规则，找没遮干净的残留。

只输出「哪一类、几处、在哪」，绝不输出命中的原文——这份自检结果本身
也要能安全地留在电脑里、贴进工单里。

用法：
    python redact_cli.py check 合同_脱敏.docx
    python redact_cli.py check 合同_脱敏.pdf --no-ocr
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field

from redact import (
    GROUP_LABELS,
    direct_runs,
    find_matches,
    iter_document_paragraphs,
    paragraph_text,
)

TEXT_EXT = {".md", ".markdown", ".txt"}
XLSX_EXT = {".xlsx", ".xlsm"}
PDF_EXT = {".pdf"}
SUPPORTED_EXT = {".docx"} | TEXT_EXT | XLSX_EXT | PDF_EXT


@dataclass(frozen=True)
class Unit:
    """一段待检查的文本，where 是给人看的位置（第几段/页/行/单元格）。"""

    where: str
    text: str


@dataclass(frozen=True)
class Finding:
    group: str
    where: str
    count: int


@dataclass
class CheckResult:
    source: str
    units: int = 0
    findings: list[Finding] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)       # 没能自检的位置（影响退出码）
    reminders: list[str] = field(default_factory=list)   # 主动跳过/格式限制（只提示）

    @property
    def total(self) -> int:
        return sum(item.count for item in self.findings)

    @property
    def clean(self) -> bool:
        return not self.findings

    @property
    def verified(self) -> bool:
        """全程都检查过了才叫 verified；有位置没查成就是 False。"""
        return not self.notes

    def by_group(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for item in self.findings:
            counts[item.group] = counts.get(item.group, 0) + item.count
        return counts


TOKEN_RE = re.compile(r"^【[^】]{1,16}】$")


def _looks_masked(value: str) -> bool:
    """掩码模式留下的「星****公司」还长得像敏感值，但已经不是原文，不算残留。"""
    return "**" in value


def _looks_like_token(value: str) -> bool:
    """条目本身就是一个代号（比如「密码：【密码】」会被密码规则再抓一次），不算残留。"""
    return bool(TOKEN_RE.match(value.strip()))


def _decode_text_file(path: str) -> str:
    raw = open(path, "rb").read()
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("gb18030")


def iter_units(
    path: str,
    *,
    ocr: bool = True,
    ocr_engine: str | None = None,
    result: CheckResult | None = None,
):
    """按格式逐段产出可检查的文本；读不了的图片页记进 result.notes。"""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        index = 0
        for paragraph, where in iter_document_paragraphs(doc):
            runs = direct_runs(paragraph)
            if not runs:
                continue
            text = paragraph_text(runs)
            if not text:
                continue
            index += 1
            yield Unit("第 %d 段（%s）" % (index, where), text)
        return

    if ext in TEXT_EXT:
        for lineno, line in enumerate(_decode_text_file(path).splitlines(), 1):
            if line.strip():
                yield Unit("第 %d 行" % lineno, line)
        return

    if ext in XLSX_EXT:
        from openpyxl import load_workbook

        wb = load_workbook(path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    where = "工作表「%s」!%s" % (ws.title, cell.coordinate)
                    if isinstance(value, str) and not value.startswith("="):
                        yield Unit(where, value)
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        # 电话/账号常被 Excel 存成数字
                        yield Unit(where, str(value))
                    comment = getattr(cell, "comment", None)
                    if comment and isinstance(comment.text, str):
                        yield Unit(where + "（批注）", comment.text)
        return

    if ext in PDF_EXT:
        import pymupdf

        doc = pymupdf.open(path)
        try:
            for page in doc:
                number = page.number + 1
                text = page.get_text("text")
                if text.strip():
                    yield Unit("第 %d 页" % number, text)
                images = page.get_images()
                if not images or not ocr:
                    if images and result is not None:
                        result.reminders.append(
                            "第 %d 页有图片，按 --no-ocr 跳过图片文字；图片里的内容没有自检" % number
                        )
                    continue
                try:
                    from formats import _ocr_normalize
                    from ocr import ocr_pdf_page

                    lines = ocr_pdf_page(page, force=ocr_engine)
                except Exception as exc:
                    if result is not None:
                        result.notes.append(
                            "第 %d 页有图片，但 OCR 不可用，图片里的文字没有自检：%s" % (number, exc)
                        )
                    continue
                joined = "\n".join(_ocr_normalize(line[0]) for line in lines)
                if joined.strip():
                    yield Unit("第 %d 页（图片 OCR）" % number, joined)
        finally:
            doc.close()
        return

    raise ValueError("不支持的格式：%s（支持 .docx / .pdf / .xlsx / .md / .txt）" % (ext or "无扩展名"))


def scan_file(path: str, *, ocr: bool = True, ocr_engine: str | None = None) -> CheckResult:
    """对一份已经脱敏的文件做自检。

    逐行匹配：规则里的「甲方：」「联系人：」这类标签一旦跨行，会把下一行的
    字段名当成公司名/人名（PDF 脱敏后标签和代号常被排版分成两行），
    那是自检自己造出来的假残留。真正的敏感值不会跨行写，逐行匹配更准。
    """
    result = CheckResult(source=os.path.abspath(path))
    for unit in iter_units(path, ocr=ocr, ocr_engine=ocr_engine, result=result):
        result.units += 1
        hits: dict[str, int] = {}
        for line in unit.text.splitlines():
            if not line.strip():
                continue
            for match in find_matches(line):
                if _looks_masked(match.value) or _looks_like_token(match.value):
                    continue
                hits[match.group] = hits.get(match.group, 0) + 1
        for group in GROUP_LABELS:
            if group in hits:
                result.findings.append(Finding(group, unit.where, hits[group]))
    return result


def format_report(result: CheckResult) -> str:
    """人看的报告：类别、数量、位置，不含原文。"""
    lines = [
        "脱敏自检（只列类别和位置，不含原文）",
        "================================",
        "文件：%s" % result.source,
        "检查：%d 处文本" % result.units,
        "",
        "命中统计：",
    ]
    counts = result.by_group()
    if counts:
        for group in GROUP_LABELS:
            if group in counts:
                lines.append("  %s：%d 处" % (GROUP_LABELS[group], counts[group]))
        lines.append(
            "  合计：%d 处 / %d 个位置" % (result.total, len(result.findings))
        )
    else:
        lines.append("  无")

    if result.findings:
        lines.append("")
        lines.append("残留清单：")
        for item in result.findings:
            lines.append(
                "  - %s：%s %d 处" % (item.where, GROUP_LABELS[item.group], item.count)
            )

    if result.reminders:
        lines.append("")
        lines.append("提示：")
        for reminder in result.reminders:
            lines.append("  - %s" % reminder)
    if result.notes:
        lines.append("")
        lines.append("未能自检：")
        for note in result.notes:
            lines.append("  - %s" % note)

    lines.append("")
    if not result.clean:
        lines.append("结果：发现 %d 处残留，请补遮后再外发" % result.total)
    elif not result.verified:
        lines.append("结果：没有发现残留，但有位置未能自检，不能算干净")
    elif result.reminders:
        lines.append("结果：没有发现残留（有内容按参数跳过，未自检，见上方提示）")
    else:
        lines.append("结果：干净")
    return "\n".join(lines)


def as_dict(result: CheckResult) -> dict:
    """机器读的结果；同样只有计数和位置。"""
    counts = result.by_group()
    return {
        "source": result.source,
        "units": result.units,
        "clean": result.clean,
        "verified": result.verified,
        "total": result.total,
        "counts": {
            group: {"category": GROUP_LABELS[group], "count": count}
            for group, count in counts.items()
        },
        "findings": [
            {
                "where": item.where,
                "group": item.group,
                "category": GROUP_LABELS[item.group],
                "count": item.count,
            }
            for item in result.findings
        ],
        "notes": list(result.notes),
        "reminders": list(result.reminders),
    }
