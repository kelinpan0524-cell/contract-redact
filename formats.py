#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""非 docx 格式的脱敏：md/txt、xlsx、pdf。

docx 的主流程在 redact.py；这里复用它的识别引擎（find_matches），
只负责「读文件 → 换成代号 → 写回去」这一段。
"""
from __future__ import annotations

import io
import json
import os
import re

import re as _re

import redact
from candidates import CandidateCollector, finalize_candidates
from redact import (
    GROUP_LABELS,
    TYPED_BASE,
    alpha_label,
    config_terms,
    find_matches,
    mask_value,
    scan_candidates,
)


_FULLWIDTH = {ord(c): ord(a) for c, a in zip(
    "０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ"
    "ａｂｃｄｅｆｇｈｉｊｋｌｍｎｏｐｑｒｓｔｕｖｗｘｙｚ：（）＠．",
    "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    "abcdefghijklmnopqrstuvwxyz:()@.",
)}


def _ocr_normalize(text: str) -> str:
    """OCR 常把全角数字、邮箱、号码拆开，匹配前先收拢。"""
    text = text.translate(_FULLWIDTH)
    text = _re.sub(r"\s*@\s*", "@", text)
    text = _re.sub(r"(?<=\d)\s+(?=\d)", "", text)
    return text

TEXT_EXT = {".md", ".markdown", ".txt"}
XLSX_EXT = {".xlsx", ".xlsm"}
PDF_EXT = {".pdf"}


class CodeBook:
    """同一实体全文复用同一代号。"""

    def __init__(self):
        self.map: dict[tuple[str, str], str] = {}
        self.used: dict[str, int] = {}
        self.seen: set[tuple[str, str]] = set()

    def token(self, group: str, value: str) -> str:
        self.seen.add((group, value))
        if redact.MASK_ONLY:
            return mask_value(value)
        key = (group, value)
        if key in self.map:
            return self.map[key]
        index = self.used.get(group, 0)
        self.used[group] = index + 1
        if group == "company":
            token = "【公司%s】" % alpha_label(index)
        elif group == "person":
            token = "【姓名%s】" % alpha_label(index)
        elif group == "customer":
            token = "【客户%s】" % alpha_label(index)
        elif group == "supplier":
            token = "【供应商%s】" % alpha_label(index)
        else:
            token = "【%s%d】" % (TYPED_BASE[group], index + 1)
        self.map[key] = token
        return token

    def mapping(self) -> dict[str, str]:
        return {token: value for (group, value), token in self.map.items()}


def _replace_all(text: str, book: CodeBook) -> str:
    matches = find_matches(text)
    if not matches:
        return text
    out = []
    cursor = 0
    for match in sorted(matches, key=lambda m: m.start):
        out.append(text[cursor:match.start])
        out.append(book.token(match.group, match.value))
        cursor = match.end
    out.append(text[cursor:])
    return "".join(out)


def _write_text(
    path: str,
    content: str,
    encoding: str = "utf-8",
    newline: str | None = "",
) -> None:
    """统一写文件：创建时就是 600，映射表、报告、候选词都可能含原文。"""
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding=encoding, newline=newline) as handle:
        handle.write(content)
    os.chmod(path, 0o600)


def _write_json(path: str, data):
    _write_text(path, json.dumps(data, ensure_ascii=False, indent=2) + "\n")


# ---------------------------------------------------------------- 纯文本


def redact_text(src: str, out: str, collector: CandidateCollector | None = None) -> CodeBook:
    raw = open(src, "rb").read()
    try:
        text = raw.decode("utf-8")
        encoding = "utf-8"
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
        encoding = "gb18030"
    for lineno, line in enumerate(text.splitlines(), 1):
        scan_candidates(line, "第 %d 行" % lineno, collector)
    book = CodeBook()
    new = _replace_all(text, book)
    _write_text(out, new, encoding=encoding, newline="")
    return book


# ---------------------------------------------------------------- Excel


def redact_xlsx(
    src: str,
    out: str,
    collector: CandidateCollector | None = None,
) -> tuple[CodeBook, list[str]]:
    from openpyxl import load_workbook

    warnings: list[str] = []
    wb = load_workbook(src)
    book = CodeBook()
    cells = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                location = "工作表「%s」!%s" % (ws.title, cell.coordinate)
                if isinstance(value, str) and not value.startswith("="):
                    scan_candidates(value, location, collector)
                    new = _replace_all(value, book)
                    if new != value:
                        cell.value = new
                        cells += 1
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    # 电话/账号常被 Excel 存成数字
                    as_text = str(value)
                    new = _replace_all(as_text, book)
                    if new != as_text:
                        cell.value = new
                        cells += 1
                comment = getattr(cell, "comment", None)
                if comment and isinstance(comment.text, str):
                    scan_candidates(comment.text, location + "（批注）", collector)
                    new_comment = _replace_all(comment.text, book)
                    if new_comment != comment.text:
                        comment.text = new_comment
                        cells += 1
        if len(ws.title) > 1:
            title = ws.title
            ws.title = _replace_all(title, book) or title
        if ws.sheet_state == "hidden":
            warnings.append("工作表「%s」是隐藏的，已一并处理" % ws.title)

    props = wb.properties
    for attr in ("creator", "lastModifiedBy", "company", "manager", "title", "subject", "description", "keywords"):
        if hasattr(props, attr):
            try:
                setattr(props, attr, "")
            except Exception:
                pass

    if cells == 0:
        warnings.append("没有单元格命中（公式不算；图片/图表里的文字不处理）")
    wb.save(out)
    return book, warnings


# ---------------------------------------------------------------- PDF


def redact_pdf(
    src: str,
    out: str,
    ocr_engine: str | None = None,
    collector: CandidateCollector | None = None,
) -> tuple[CodeBook, list[str]]:
    import pymupdf

    warnings: list[str] = []
    doc = pymupdf.open(src)
    book = CodeBook()

    # 第一遍：按页收集，先建好全局代号（同一实体跨页同码）
    page_data = []
    for page in doc:
        text = page.get_text("text")
        scan_candidates(text, "第 %d 页" % (page.number + 1), collector)
        page_data.append((page, text, find_matches(text)))
    for _, _text, matches in page_data:
        for match in matches:
            book.token(match.group, match.value)

    replaced = 0
    ocr_pages = 0
    for page, text, matches in page_data:
        if not text.strip():
            if not page.get_images():
                continue
            # 扫描件：走系统 OCR，按行重写
            try:
                from ocr import ocr_pdf_page
            except Exception:
                warnings.append("第 %d 页是扫描件，但本机没有可用的 OCR 引擎" % (page.number + 1))
                continue
            try:
                lines = ocr_pdf_page(page, force=ocr_engine)
            except Exception as exc:
                warnings.append("第 %d 页 OCR 失败：%s" % (page.number + 1, exc))
                continue
            ocr_pages += 1
            rect = page.rect
            pending = []
            for line_text, bx, by, bw, bh in lines:
                normalized = _ocr_normalize(line_text)
                scan_candidates(
                    normalized,
                    "第 %d 页（扫描件）" % (page.number + 1),
                    collector,
                )
                line_matches = find_matches(normalized)
                if not line_matches:
                    continue
                new_text = []
                cursor = 0
                for match in line_matches:
                    new_text.append(normalized[cursor:match.start])
                    new_text.append(book.token(match.group, match.value))
                    cursor = match.end
                new_text.append(normalized[cursor:])
                rewritten = "".join(new_text)
                box = pymupdf.Rect(
                    rect.x0 + bx * rect.width,
                    rect.y0 + (1 - (by + bh)) * rect.height,
                    rect.x0 + (bx + bw) * rect.width,
                    rect.y0 + (1 - by) * rect.height,
                )
                page.add_redact_annot(box, fill=(1, 1, 1))
                pending.append((box, rewritten))
                replaced += 1
            try:
                page.apply_redactions()
            except Exception as exc:
                warnings.append("第 %d 页擦除失败：%s" % (page.number + 1, exc))
            for box, rewritten in pending:
                page.insert_text(
                    (box.x0 + 1, box.y1 - 2),
                    rewritten,
                    fontname="china-s",
                    fontsize=max(6, min(14, box.height * 0.8)),
                    color=(0, 0, 0),
                )
            continue
        seen = set()
        pending = []
        for match in matches:
            key = (match.group, match.value)
            if key in seen:
                continue
            seen.add(key)
            token = book.token(*key)
            for rect in page.search_for(match.value):
                page.add_redact_annot(rect, fill=(1, 1, 1))
                pending.append((rect, token))
        try:
            page.apply_redactions()   # 先把原文字真正抹掉
        except Exception as exc:      # 老版本参数差异
            warnings.append("第 %d 页擦除失败：%s" % (page.number + 1, exc))
        for rect, token in pending:   # 再写代号，避免被一并擦掉
            page.insert_text(
                (rect.x0, rect.y1 - 2.5),
                token,
                fontname="china-s",
                fontsize=max(7, min(12, rect.height * 0.75)),
                color=(0, 0, 0),
            )
            replaced += 1

    try:
        doc.set_metadata({})
    except Exception:
        pass
    if ocr_pages:
        warnings.append("有 %d 页是扫描件，已用系统 OCR 识别后按行重写；排版会变、识别可能有错字，务必肉眼核对" % ocr_pages)
    if replaced == 0:
        warnings.append("PDF 里没有命中任何实体")
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return book, warnings


# ---------------------------------------------------------------- 统一入口


def redact_any(
    src: str,
    out: str,
    map_path: str,
    report_path: str | None,
    cfg_path: str | None,
    ocr_engine: str | None = None,
    candidates_path: str | None = None,
) -> int:
    """md/txt、xlsx、pdf 的入口；返回 0 表示成功。"""
    ext = os.path.splitext(src)[1].lower()
    warnings: list[str] = []
    collector = CandidateCollector(config_terms()) if candidates_path else None

    if ext in TEXT_EXT:
        book = redact_text(src, out, collector)
    elif ext in XLSX_EXT:
        book, warnings = redact_xlsx(src, out, collector)
    elif ext in PDF_EXT:
        book, warnings = redact_pdf(src, out, ocr_engine, collector)
    else:
        print("错误：暂不支持 %s（支持 .docx .doc .pdf .md .txt .xlsx）" % ext)
        return 2

    mapping = {} if redact.MASK_ONLY else book.mapping()
    if mapping:
        _write_json(map_path, mapping)

    lines = [
        "脱敏报告",
        "========",
        "源文件：%s" % src,
        "输出文件：%s" % out,
        "映射表：%s（权限 600，内含原文，切勿外发）" % map_path,
        "命中统计：",
    ]
    counts: dict[str, int] = {}
    for (group, _value) in book.seen:
        counts[group] = counts.get(group, 0) + 1
    for group, count in sorted(counts.items()):
        lines.append("  %s：%d 个实体" % (GROUP_LABELS.get(group, group), count))
    if not counts:
        lines.append("  无")
    candidate_summary, candidate_reminder = finalize_candidates(
        collector, candidates_path, src,
    )
    lines.append("候选词：%s" % candidate_summary)
    if candidate_reminder:
        lines.append("  （%s）" % candidate_reminder)
    lines.append("")
    lines.append("提醒：")
    if warnings:
        lines.extend("  - " + w for w in warnings)
    lines.append("  - 规则识别不保证遮干净，外发前请人工复核")
    text = "\n".join(lines)
    print(text)
    if report_path:
        _write_text(report_path, text + "\n")
    return 0


# ---------------------------------------------------------------- 还原


def _restore_strings(text: str, mapping: dict[str, str]) -> str:
    for token, original in sorted(mapping.items(), key=lambda kv: -len(kv[0])):
        if token in text:
            text = text.replace(token, original)
    return text


def restore_text(src: str, out: str, mapping: dict[str, str]) -> int:
    raw = open(src, "rb").read()
    try:
        text = raw.decode("utf-8")
        encoding = "utf-8"
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
        encoding = "gb18030"
    _write_text(out, _restore_strings(text, mapping), encoding=encoding, newline="")
    return 0


def restore_xlsx(src: str, out: str, mapping: dict[str, str]) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(src)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    cell.value = _restore_strings(cell.value, mapping)
                comment = getattr(cell, "comment", None)
                if comment and isinstance(comment.text, str):
                    comment.text = _restore_strings(comment.text, mapping)
    wb.save(out)
    return 0


def restore_pdf(src: str, out: str, mapping: dict[str, str]) -> int:
    """把代号换回原文。按 span 处理，避免字体渲染出的字距导致 search_for 找不到。"""
    import pymupdf

    doc = pymupdf.open(src)
    for page in doc:
        pending = []
        data = page.get_text("dict")
        for block in data.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    if not text.strip():
                        continue
                    compact = _re.sub(r"\s+", "", text)
                    restored = _restore_strings(compact, mapping)
                    if restored == compact:
                        continue
                    rect = pymupdf.Rect(span["bbox"])
                    page.add_redact_annot(rect, fill=(1, 1, 1))
                    pending.append((rect, restored, span.get("size", 10)))
        if not pending:
            continue
        try:
            page.apply_redactions()
        except Exception:
            continue
        for rect, text, size in pending:
            page.insert_text(
                (rect.x0 + 1, rect.y1 - 2),
                text,
                fontname="china-s",
                fontsize=max(6, min(14, size)),
                color=(0, 0, 0),
            )
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return 0


def restore_any(src: str, out: str, mapping: dict[str, str]) -> int:
    ext = os.path.splitext(src)[1].lower()
    if ext in TEXT_EXT:
        return restore_text(src, out, mapping)
    if ext in XLSX_EXT:
        return restore_xlsx(src, out, mapping)
    if ext in PDF_EXT:
        return restore_pdf(src, out, mapping)
    print("错误：暂不支持还原 %s" % ext)
    return 2
