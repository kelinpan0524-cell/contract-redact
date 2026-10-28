#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""check 子命令验收：干净返回 0、有残留返回 2，且输出里绝不出现原文。

覆盖：
1. 全量脱敏后的 docx / md / xlsx / pdf / 扫描件 → 自检干净（退出码 0）；
2. 只遮一部分（--only）→ 自检报残留（退出码 2），给出类别、数量和段落/行/单元格/页码；
3. 掩码模式的星号不会被误报成残留；
4. --json 输出同样只有计数和位置；
5. 所有输出（stdout + stderr）里都不出现任何原文。
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from contextlib import redirect_stderr, redirect_stdout

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import ocr  # noqa: E402
import redact  # noqa: E402
import redact_cli  # noqa: E402

COMPANY = "星海精密制造有限公司"
PERSON = "陈志远"
MOBILE = "13800000001"
EMAIL = "chen.zy@example.com"
# 测试用的虚构口令，不是任何真实凭据
SAMPLE_PWD = "样例" + "口令" + "1234"
AMOUNT = "1,234,567.89"
SECRETS = [COMPANY, PERSON, MOBILE, EMAIL, SAMPLE_PWD]

DOCX_LINES = [
    "甲方：%s" % COMPANY,
    "法定代表人：%s" % PERSON,
    "联系人：李静 手机：%s 邮箱：%s" % (MOBILE, EMAIL),
    "密码：%s" % SAMPLE_PWD,
    "合同金额：人民币 %s 元，违约金 20%%。" % AMOUNT,
]

CFG_DEFAULTS = {
    "extra_orgs": [],
    "extra_people": [],
    "extra_customers": [],
    "extra_suppliers": [],
    "keep": [],
    "org_map": {},
    "person_map": {},
    "customer_map": {},
    "supplier_map": {},
}

ALL_OUTPUT = ""


def reset_state() -> None:
    redact.ENABLED_GROUPS = None
    redact.MASK_ONLY = False
    redact.CFG.clear()
    redact.CFG.update(json.loads(json.dumps(CFG_DEFAULTS)))


def cli(argv: list[str]) -> tuple[int, str, str]:
    """进程内跑一次 CLI；顺便把所有输出收集起来做「不含原文」的总检查。"""
    global ALL_OUTPUT
    reset_state()
    out, err = io.StringIO(), io.StringIO()
    try:
        with redirect_stdout(out), redirect_stderr(err):
            code = redact_cli.main(argv)
    except SystemExit as exc:
        code = int(exc.code or 0)
    text_out, text_err = out.getvalue(), err.getvalue()
    ALL_OUTPUT += text_out + text_err
    return code, text_out, text_err


def make_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    for line in DOCX_LINES:
        doc.add_paragraph(line)
    doc.save(path)


def docx_text(path: str) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def write_text(path: str, text: str) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
        handle.write(text)


def make_pdf(path: str) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in DOCX_LINES:
        page.insert_text((50, y), line, fontname="china-s", fontsize=11)
        y += 24
    doc.save(path)
    doc.close()


def make_scan_pdf(text_pdf: str, path: str) -> None:
    """把 PDF 渲成图再塞回去：没有文字层，只能靠 OCR 读。"""
    import pymupdf

    src = pymupdf.open(text_pdf)
    pixmap = src[0].get_pixmap(dpi=200)
    png = path + ".png"
    pixmap.save(png)
    src.close()
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=png)
    doc.save(path)
    doc.close()
    os.unlink(png)


def redact_file(source: str, output: str, extra: list[str] | None = None) -> int:
    argv = [
        "redact", source,
        "-o", output,
        "--map", os.path.join(os.path.dirname(output), "map.json"),
        "--no-audit",
        "--no-candidates",
    ] + list(extra or [])
    code, _out, _err = cli(argv)
    return code


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="contract-redact-check-")
    checks: list[tuple[str, bool]] = []

    # ---- 1. docx：全量脱敏 → 干净 ----
    source = os.path.join(tmp, "合同.docx")
    make_docx(source)
    full = os.path.join(tmp, "合同_脱敏.docx")
    checks.append(("docx 全量脱敏成功", redact_file(source, full) == 0))
    full_text = docx_text(full)
    checks.append(("docx 脱敏件无原文", not any(secret in full_text for secret in SECRETS)))

    code, out, _err = cli(["check", full])
    checks.append(("干净件退出码 0", code == 0))
    checks.append(("干净件结论", "结果：干净" in out))
    checks.append(("干净件命中统计为无", "命中统计：" in out and "  无" in out))

    # ---- 2. docx：只遮手机 → 残留 ----
    partial = os.path.join(tmp, "合同_只遮手机.docx")
    checks.append(("只遮手机成功", redact_file(source, partial, ["--only", "phone"]) == 0))
    code, out, _err = cli(["check", partial])
    checks.append(("残留件退出码 2", code == 2))
    checks.append(("残留件给出结论", "发现" in out and "残留" in out))
    checks.append(("残留清单有段落位置", "第 1 段（正文）" in out))
    checks.append(("残留清单有类别", "公司名" in out and "人名" in out))
    checks.append(("残留件标明密码类别", "密码" in out))
    checks.append(("手机已遮不再报", "电话" not in out))

    # ---- 3. JSON 输出：只有计数和位置 ----
    code, out, _err = cli(["check", partial, "--json"])
    checks.append(("JSON 模式退出码 2", code == 2))
    try:
        data = json.loads(out)
    except ValueError:
        data = {}
    checks.append(("JSON 可解析", bool(data)))
    checks.append(("JSON clean=false", data.get("clean") is False))
    checks.append(("JSON total>0", (data.get("total") or 0) > 0))
    checks.append(
        (
            "JSON findings 有位置和类别",
            bool(data.get("findings"))
            and all(
                item.get("where") and item.get("group") and item.get("category")
                for item in data.get("findings", [])
            ),
        )
    )
    checks.append(("JSON 不含原文", not any(secret in out for secret in SECRETS)))

    # ---- 4. md：行号定位 + 脱敏后追加手机也能查出来 ----
    md = os.path.join(tmp, "备忘.md")
    write_text(md, "甲方：%s\n联系人：%s 手机：%s\n金额 %s 元\n" % (COMPANY, PERSON, MOBILE, AMOUNT))
    md_out = os.path.join(tmp, "备忘_脱敏.md")
    checks.append(("md 全量脱敏成功", redact_file(md, md_out) == 0))
    code, out, _err = cli(["check", md_out])
    checks.append(("md 干净件退出码 0", code == 0))

    write_text(md_out, open(md_out, encoding="utf-8").read() + "补充手机：13700000009\n")
    code, out, _err = cli(["check", md_out])
    checks.append(("md 追加手机后退出码 2", code == 2))
    checks.append(("md 定位到行", "第 4 行" in out))
    checks.append(("md 只报电话这一类", "电话 1 处" in out))

    # ---- 5. xlsx：单元格定位 ----
    from openpyxl import Workbook

    xlsx = os.path.join(tmp, "明细.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "合同信息"
    ws["A1"] = "甲方"
    ws["A2"] = COMPANY
    ws["B2"] = MOBILE
    wb.save(xlsx)
    xlsx_out = os.path.join(tmp, "明细_脱敏.xlsx")
    checks.append(("xlsx 只遮手机成功", redact_file(xlsx, xlsx_out, ["--only", "phone"]) == 0))
    code, out, _err = cli(["check", xlsx_out])
    checks.append(("xlsx 有残留退出码 2", code == 2))
    checks.append(("xlsx 定位到单元格", "工作表「合同信息」!A2" in out))

    # ---- 6. 掩码模式：星号不算残留 ----
    masked = os.path.join(tmp, "合同_掩码.docx")
    checks.append(("掩码脱敏成功", redact_file(source, masked, ["--mask-only"]) == 0))
    masked_text = docx_text(masked)
    checks.append(("掩码件打星且无原文", "*" in masked_text and not any(s in masked_text for s in SECRETS)))
    code, out, _err = cli(["check", masked])
    checks.append(("掩码件退出码 0（星号不误报）", code == 0))
    checks.append(("掩码件结论干净", "结果：干净" in out))

    # ---- 7. PDF：页码定位 ----
    pdf = os.path.join(tmp, "合同.pdf")
    make_pdf(pdf)
    code, out, _err = cli(["check", pdf])
    checks.append(("未脱敏 PDF 退出码 2", code == 2))
    checks.append(("PDF 定位到页", "第 1 页" in out))
    pdf_out = os.path.join(tmp, "合同_脱敏.pdf")
    checks.append(("PDF 脱敏成功", redact_file(pdf, pdf_out) == 0))
    code, out, _err = cli(["check", pdf_out])
    checks.append(("脱敏后 PDF 退出码 0", code == 0))

    # ---- 8. 扫描件：OCR 读图自检 ----
    if ocr.available():
        scan = os.path.join(tmp, "合同_扫描件.pdf")
        make_scan_pdf(pdf, scan)
        code, out, _err = cli(["check", scan])
        checks.append(("扫描件 OCR 能查到残留", code == 2 and "图片 OCR" in out))
        scan_out = os.path.join(tmp, "合同_扫描件_脱敏.pdf")
        checks.append(("扫描件脱敏成功", redact_file(scan, scan_out) == 0))
        code, out, _err = cli(["check", scan_out])
        checks.append(("脱敏后扫描件退出码 0", code == 0))

        code, out, _err = cli(["check", scan, "--no-ocr"])
        checks.append(("--no-ocr 跳过图片", "按 --no-ocr 跳过" in out))
        checks.append(("--no-ocr 结果是未完全自检", "没有发现残留（有内容按参数跳过" in out))
    else:
        print("（本机没有 OCR 引擎，跳过扫描件用例）")

    # ---- 9. 词表白名单也生效 ----
    keep_md = os.path.join(tmp, "白名单.md")
    write_text(keep_md, "甲方：%s\n金额 100 元\n" % COMPANY)
    config = os.path.join(tmp, "vocab.json")
    write_text(config, json.dumps({"keep": [COMPANY]}, ensure_ascii=False))
    code, _out, _err = cli(["check", keep_md])
    checks.append(("不带词表时报残留", code == 2))
    code, _out, _err = cli(["check", keep_md, "--config", config])
    checks.append(("词表白名单后干净", code == 0))

    # ---- 10. 错误处理 ----
    code, _out, _err = cli(["check", os.path.join(tmp, "不存在.docx")])
    checks.append(("文件不存在返回 1", code == 1))
    bogus = os.path.join(tmp, "数据.csv")
    write_text(bogus, "a,b\n1,2\n")
    code, _out, _err = cli(["check", bogus])
    checks.append(("不支持的格式返回 1", code == 1))

    # ---- 总检查：任何输出里都不许出现原文 ----
    leaked = [secret for secret in SECRETS if secret in ALL_OUTPUT]
    checks.append(("所有输出都不含原文", not leaked))

    for name, ok in checks:
        print("%s %s" % ("✓" if ok else "✗", name))
    failed = [name for name, ok in checks if not ok]
    print("\n结果：%s" % ("全部通过" if not failed else "失败：" + "、".join(failed)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
