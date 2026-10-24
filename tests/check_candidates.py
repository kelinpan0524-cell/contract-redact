#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""候选词闭环验收：发现 → 预览 → 合并 → 再脱敏。

验证五件事：
1. 词表外的机构名/人名会进候选文件（权限 600），词表里已有的不重复出现；
2. merge_candidates.py 的 --dry-run 不写文件，--only 能按分类筛选；
3. 合并只增不减：keep / aliases / 客户表 / 手写字段都还在，重复跑不重复；
4. 合并后再脱敏，这些名字被正常替换，候选文件不再生成；
5. md 的候选位置记的是行号，带「供应商：」标签的归到 supplier。
"""
from __future__ import annotations

import io
import json
import os
import stat
import sys
import tempfile
from contextlib import redirect_stdout

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))
sys.path.insert(0, HERE)

import merge_candidates  # noqa: E402
import redact  # noqa: E402

UNLISTED_ORG = "青禾远航精密制造有限公司"
UNLISTED_PERSON = "赵启明"
LISTED_ORG = "北辰智联科技有限公司"
LISTED_PERSON = "钱多多"
AMOUNT = "888,888.00"


def call_main(module, argv: list[str]) -> str:
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        code = module.main(argv)
    if code != 0:
        raise RuntimeError("%s 返回 %s" % (module.__name__, code))
    return buffer.getvalue()


def write_text(path: str, text: str) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
        handle.write(text)


def read_text(path: str) -> str:
    with open(path, encoding="utf-8") as handle:
        return handle.read()


def make_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("甲方：%s" % UNLISTED_ORG)
    doc.add_paragraph("法定代表人：%s" % UNLISTED_PERSON)
    doc.add_paragraph("乙方：%s" % LISTED_ORG)
    doc.add_paragraph("联系人：%s" % LISTED_PERSON)
    doc.add_paragraph("合同金额：人民币 %s 元" % AMOUNT)
    doc.save(path)


def docx_text(path: str) -> str:
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def mode_of(path: str) -> int:
    return stat.S_IMODE(os.stat(path).st_mode)


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="contract-redact-candidates-")
    config = os.path.join(tmp, "vocab.json")
    source = os.path.join(tmp, "合同.docx")
    redacted = os.path.join(tmp, "合同_脱敏.docx")
    candidates = os.path.join(tmp, "词表外候选.json")

    original_config = {
        "extra_orgs": [LISTED_ORG],
        "extra_people": [LISTED_PERSON],
        "extra_customers": ["ACME"],
        "keep": ["市仲裁委员会"],
        "aliases": {LISTED_ORG: ["北辰"]},
        "手写备注": "不要丢",
    }
    write_text(config, json.dumps(original_config, ensure_ascii=False, indent=2))
    os.chmod(config, 0o600)
    make_docx(source)

    checks: list[tuple[str, bool]] = []

    # ---- 0. 后缀套后缀只报完整名称 ----
    nested = redact.discover_candidates("受托方：青禾建筑设计院有限公司")
    nested_orgs = [value for category, value in nested if category == "org"]
    checks.append(("嵌套后缀只报全名", nested_orgs == ["青禾建筑设计院有限公司"]))

    # ---- 1. 脱敏时发现候选 ----
    output = call_main(
        redact,
        [
            source,
            "-o", redacted,
            "--map", os.path.join(tmp, "map.json"),
            "--config", config,
            "--candidates", candidates,
        ],
    )
    data = json.loads(read_text(candidates))
    items = {(item["category"], item["text"]): item for item in data["candidates"]}
    redacted_text = docx_text(redacted)

    checks.append(("候选文件权限 600", mode_of(candidates) == 0o600))
    checks.append(("发现词表外机构", ("org", UNLISTED_ORG) in items))
    checks.append(("发现词表外人名", ("people", UNLISTED_PERSON) in items))
    checks.append(("词表内机构不重复", ("org", LISTED_ORG) not in items))
    checks.append(("词表内人名不重复", ("people", LISTED_PERSON) not in items))
    checks.append(("候选总数正确", data["summary"]["total"] == 2))
    checks.append(
        ("首次位置是段落", items[("org", UNLISTED_ORG)]["first_seen"].startswith("第 1 段"))
    )
    checks.append(("终端打印候选摘要", "发现 2 个词表外实体" in output))
    checks.append(("脱敏件无原名", UNLISTED_ORG not in redacted_text))
    checks.append(("金额未动", AMOUNT in redacted_text))

    # ---- 2. dry-run + --only ----
    before = read_text(config)
    dry = call_main(
        merge_candidates,
        [candidates, "--config", config, "--only", "people", "--dry-run"],
    )
    after = read_text(config)
    checks.append(("dry-run 不写文件", before == after))
    checks.append(
        ("--only 只列人名", UNLISTED_PERSON in dry and UNLISTED_ORG not in dry)
    )

    # ---- 3. 按分类合并 ----
    call_main(merge_candidates, [candidates, "--config", config, "--only", "org"])
    merged = json.loads(read_text(config))
    checks.append(
        (
            "只并入机构",
            UNLISTED_ORG in merged["extra_orgs"]
            and UNLISTED_PERSON not in merged["extra_people"],
        )
    )

    second = call_main(merge_candidates, [candidates, "--config", config])
    merged = json.loads(read_text(config))
    checks.append(("人名随后并入", UNLISTED_PERSON in merged["extra_people"]))
    checks.append(("重复项被跳过", "已在 extra_orgs" in second))
    checks.append(("keep 保留", merged.get("keep") == ["市仲裁委员会"]))
    checks.append(("aliases 保留", merged.get("aliases") == {LISTED_ORG: ["北辰"]}))
    checks.append(("客户表保留", merged.get("extra_customers") == ["ACME"]))
    checks.append(("手写字段保留", merged.get("手写备注") == "不要丢"))
    checks.append(("词表权限 600", mode_of(config) == 0o600))

    third = call_main(merge_candidates, [candidates, "--config", config])
    merged = json.loads(read_text(config))
    checks.append(("重复合并新增 0 条", "将并入 0 条" in third and "没有新词" in third))
    checks.append(
        (
            "没有重复项",
            len(merged["extra_orgs"]) == 2 and len(merged["extra_people"]) == 2,
        )
    )

    # ---- 4. 再脱敏：不再产生候选，实体被正常替换 ----
    redacted2 = os.path.join(tmp, "合同_脱敏2.docx")
    candidates2 = os.path.splitext(redacted2)[0] + ".candidates.json"
    output2 = call_main(
        redact,
        [
            source,
            "-o", redacted2,
            "--map", os.path.join(tmp, "map2.json"),
            "--config", config,
        ],
    )
    text2 = docx_text(redacted2)
    checks.append(("默认候选路径未生成文件", not os.path.exists(candidates2)))
    checks.append(("报告写明未发现候选", "未发现词表外实体" in output2))
    checks.append(
        ("并入后正常替换", UNLISTED_ORG not in text2 and UNLISTED_PERSON not in text2)
    )
    checks.append(("代号仍然存在", "【公司" in text2 and "【姓名" in text2))
    checks.append(("再跑金额仍未动", AMOUNT in text2))

    # ---- 5. md：行号位置 + 供应商标签分类 ----
    md = os.path.join(tmp, "备忘.md")
    write_text(md, "供应商：蓝谷新材料有限公司\n联系人：孙立\n")
    md_out = os.path.join(tmp, "备忘_脱敏.md")
    call_main(
        redact,
        [md, "-o", md_out, "--map", os.path.join(tmp, "md.map.json"), "--config", config],
    )
    md_candidates = os.path.splitext(md_out)[0] + ".candidates.json"
    md_data = json.loads(read_text(md_candidates))
    md_items = {item["text"]: item for item in md_data["candidates"]}
    checks.append(("md 默认候选路径", os.path.isfile(md_candidates)))
    checks.append(("md 行号定位", md_items["蓝谷新材料有限公司"]["first_seen"] == "第 1 行"))
    checks.append(("供应商标签分类", md_items["蓝谷新材料有限公司"]["category"] == "supplier"))
    checks.append(("md 人名定位", md_items["孙立"]["first_seen"] == "第 2 行"))
    checks.append(("md 脱敏生效", "蓝谷新材料有限公司" not in read_text(md_out)))

    for name, ok in checks:
        print("%s %s" % ("✓" if ok else "✗", name))
    failed = [name for name, ok in checks if not ok]
    print("\n结果：%s" % ("全部通过" if not failed else "失败：" + "、".join(failed)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
