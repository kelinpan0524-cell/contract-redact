#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成多格式虚构样例：md / xlsx / pdf / doc（全部虚构）。"""
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))

LINES = [
    "采购框架合同",
    "甲方：星海精密制造有限公司",
    "法定代表人：陈志远",
    "联系人电话：13800000001",
    "邮箱：chen.zy@example.com",
    "地址：江苏省苏州市工业园区星海路88号",
    "统一社会信用代码：91340200MA2T8K9L3D",
    "合同金额 1,234,567.89 元，违约金 20%。",
]

SECRETS = ["星海精密制造有限公司", "陈志远", "13800000001", "chen.zy@example.com", "星海路88号",
           "91340200MA2T8K9L3D"]


def _write_text(path, text):
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
        handle.write(text)


def make_md():
    path = os.path.join(HERE, "sample.md")
    _write_text(path, "# 采购框架合同\n\n" + "\n".join("- " + l for l in LINES[1:]) + "\n")
    return path


def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    wb = Workbook()
    ws = wb.active
    ws.title = "合同信息"
    for i, line in enumerate(LINES, 1):
        ws.cell(row=i, column=1, value=line)
    ws.cell(row=3, column=2, value=13800000001)
    ws["B3"].comment = Comment("对接人陈志远，手机 13800000001", "法务")
    wb.properties.creator = "陈志远"
    path = os.path.join(HERE, "sample.xlsx")
    wb.save(path)
    return path


def make_pdf():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 80
    for line in LINES:
        page.insert_text((60, y), line, fontname="china-s", fontsize=11)
        y += 24
    path = os.path.join(HERE, "sample.pdf")
    doc.save(path)
    doc.close()
    return path


def make_scan_pdf():
    """把 PDF 页面渲成图再塞回 PDF：没有文字层，模拟扫描件。"""
    import pymupdf

    text_pdf = os.path.join(HERE, "sample.pdf")
    if not os.path.isfile(text_pdf):
        make_pdf()
    src = pymupdf.open(text_pdf)
    pix = src[0].get_pixmap(dpi=200)
    png = os.path.join(HERE, "_scan_tmp.png")
    pix.save(png)
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=png)
    path = os.path.join(HERE, "sample_scan.pdf")
    doc.save(path)
    doc.close()
    os.unlink(png)
    return path


def make_doc():
    from docx import Document

    d = Document()
    for line in LINES:
        d.add_paragraph(line)
    tmp = os.path.join(HERE, "sample.docx")
    d.save(tmp)
    path = os.path.join(HERE, "sample.doc")
    subprocess.run(["textutil", "-convert", "doc", "-output", path, tmp], check=True)
    return path


def main():
    paths = []
    for fn in (make_md, make_xlsx, make_pdf, make_scan_pdf):
        paths.append(fn())
    try:
        paths.append(make_doc())
    except Exception as exc:
        print("跳过 .doc（本机 textutil 不可用）：%s" % exc)
    for p in paths:
        print("已生成", p)
    print("敏感串：", "、".join(SECRETS))
    return 0


if __name__ == "__main__":
    sys.exit(main())
