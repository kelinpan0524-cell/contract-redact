#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""统一入口验收：redact / restore / check / serve 四个子命令。

验证：
1. 子命令路由、--help、不认识的子命令；
2. redact 子命令内部复用 redact.py：同样的输入产生同样的脱敏文本；
3. redact 结束打印一行 hits 摘要（各类别替换处数）；
4. --audit-json 追加一行 JSON 审计记录，权限 600，只有元数据和计数、没有原文；
5. restore 子命令逐字还原；
6. 原有参数（--only / --mask-only）照常透传。

为了不依赖 shell，测试直接调用 redact_cli.main()；每次调用前把 redact 的
模块级开关复位，模拟一次全新的进程（这正是命令行每次启动时的状态）。
"""
from __future__ import annotations

import io
import json
import os
import stat
import sys
import tempfile
from contextlib import redirect_stderr, redirect_stdout

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import redact  # noqa: E402
import redact_cli  # noqa: E402
import service  # noqa: E402
from redact import direct_runs, iter_document_paragraphs, paragraph_text  # noqa: E402

COMPANY = "青禾远航精密制造有限公司"
PERSON = "赵启明"
MOBILE = "13912345678"
AMOUNT = "777,777.00"
SECRETS = [COMPANY, PERSON, MOBILE]

CFG_DEFAULTS = {
    "extra_orgs": [],
    "extra_people": [],
    "extra_customers": [],
    "extra_suppliers": [],
    "keep": [],
    "org_map": {},
    "person_map": {},
    "customer_map": {},
    "supplier_map": {},
}


def reset_state() -> None:
    """模拟新进程：清掉上一次调用留下的开关和词表。"""
    redact.ENABLED_GROUPS = None
    redact.MASK_ONLY = False
    redact.CFG.clear()
    redact.CFG.update(json.loads(json.dumps(CFG_DEFAULTS)))


def cli(argv: list[str]) -> tuple[int, str, str]:
    """在进程内跑一次 CLI，返回（退出码，stdout，stderr）。"""
    reset_state()
    out, err = io.StringIO(), io.StringIO()
    try:
        with redirect_stdout(out), redirect_stderr(err):
            code = redact_cli.main(argv)
    except SystemExit as exc:  # 子命令解析失败时 argparse 会抛 SystemExit
        code = int(exc.code or 0)
    return code, out.getvalue(), err.getvalue()


def make_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("甲方：%s" % COMPANY)
    doc.add_paragraph("法定代表人：%s" % PERSON)
    doc.add_paragraph("手机：%s" % MOBILE)
    doc.add_paragraph("合同金额：人民币 %s 元" % AMOUNT)
    doc.save(path)


def docx_text(path: str) -> str:
    from docx import Document

    doc = Document(path)
    return "\n".join(
        paragraph_text(direct_runs(paragraph))
        for paragraph, _where in iter_document_paragraphs(doc)
    )


def mode_of(path: str) -> int:
    return stat.S_IMODE(os.stat(path).st_mode)


def write_text(path: str, text: str) -> None:
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
        handle.write(text)


def read_text(path: str) -> str:
    descriptor = os.open(path, os.O_RDONLY)
    with os.fdopen(descriptor, "r", encoding="utf-8") as handle:
        return handle.read()


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="contract-redact-cli-")
    source = os.path.join(tmp, "合同.docx")
    make_docx(source)
    original = docx_text(source)

    checks: list[tuple[str, bool]] = []

    # ---- 1. 入口与路由 ----
    code, out, _err = cli(["--help"])
    checks.append(("顶层 --help 返回 0", code == 0))
    checks.append(
        (
            "四个子命令都列出来",
            all(name in out for name in ("redact", "restore", "check", "serve")),
        )
    )
    code, _out, err = cli(["rm-all"])
    checks.append(("不认识的子命令返回 2", code == 2))
    checks.append(("提示不认识的子命令", "不认识的子命令" in err))

    code, out, _err = cli(["redact", "--help"])
    checks.append(("redact --help 返回 0", code == 0))
    checks.append(("redact --help 复用原参数说明", "输入的 .docx 文件" in out))

    checks.append(("默认监听 127.0.0.1", service.DEFAULT_HOST == "127.0.0.1"))
    checks.append(("默认端口 8791", service.DEFAULT_PORT == 8791))
    code, out, _err = cli(["serve", "--help"])
    checks.append(
        (
            "serve --help 列出关键参数",
            code == 0 and "--port" in out and "--session-file" in out,
        )
    )

    # ---- 2. redact 子命令 ----
    redacted = os.path.join(tmp, "合同_脱敏.docx")
    map_path = os.path.join(tmp, "合同_脱敏.map.json")
    report_path = os.path.join(tmp, "合同_脱敏.report.txt")
    audit_path = os.path.join(tmp, "audit.jsonl")
    code, out, err = cli(
        [
            "redact", source,
            "-o", redacted,
            "--map", map_path,
            "--report", report_path,
            "--audit-json", audit_path,
            "--no-audit",
            "--no-candidates",
        ]
    )
    checks.append(("redact 子命令返回 0", code == 0))
    checks.append(("redact 没有报错输出", err == ""))
    redacted_text = docx_text(redacted)
    checks.append(("脱敏件没有原文", not any(secret in redacted_text for secret in SECRETS)))
    checks.append(("脱敏件有代号", "【公司A】" in redacted_text and "【姓名A】" in redacted_text))
    checks.append(("金额未动", AMOUNT in redacted_text))
    checks.append(("映射表权限 600", mode_of(map_path) == 0o600))

    hits_lines = [line for line in out.splitlines() if line.startswith("hits:")]
    checks.append(("结束时打印一行 hits", len(hits_lines) == 1))
    checks.append(
        (
            "hits 里有类别和数量",
            bool(hits_lines)
            and "公司名 1" in hits_lines[0]
            and "人名 1" in hits_lines[0]
            and "电话 1" in hits_lines[0],
        )
    )
    checks.append(("hits 合计 3 处", bool(hits_lines) and "合计 3 处" in hits_lines[0]))
    checks.append(("报告文件已生成", os.path.isfile(report_path)))

    # 审计 JSON：一行、只含计数、权限 600
    checks.append(("审计 JSON 已生成", os.path.isfile(audit_path)))
    if os.path.isfile(audit_path):
        checks.append(("审计 JSON 权限 600", mode_of(audit_path) == 0o600))
        raw_audit = open(audit_path, encoding="utf-8").read()
        lines = [line for line in raw_audit.splitlines() if line.strip()]
        checks.append(("审计 JSON 一行一次运行", len(lines) == 1))
        entry = json.loads(lines[0]) if lines else {}
        checks.append(("审计 mode 正确", entry.get("mode") == "reversible"))
        checks.append(("审计 hits 正确", (entry.get("hits") or {}).get("phone") == 1))
        checks.append(
            (
                "审计 source/output 是用户文件",
                entry.get("source") == os.path.abspath(source)
                and entry.get("output") == os.path.abspath(redacted),
            )
        )
        checks.append(("审计有映射指纹", len(entry.get("mapping_sha256") or "") == 16))
        checks.append(("审计 JSON 没有原文", not any(secret in raw_audit for secret in SECRETS)))

    # 内部复用：直接调用 redact.main 应该得到同样的脱敏文本
    direct_out = os.path.join(tmp, "direct.docx")
    reset_state()
    with redirect_stdout(io.StringIO()):
        direct_code = redact.main(
            [
                source,
                "-o", direct_out,
                "--map", os.path.join(tmp, "direct.map.json"),
                "--no-audit",
                "--no-candidates",
            ]
        )
    checks.append(("redact.py 直接跑成功", direct_code == 0))
    checks.append(("子命令与 redact.py 结果一致", docx_text(direct_out) == redacted_text))

    # 非 docx 走 formats.py，也要打印 hits 摘要
    md_source = os.path.join(tmp, "备忘.md")
    write_text(
        md_source,
        "甲方：%s\n手机：%s\n合同金额 %s 元，违约金 20%%\n" % (COMPANY, MOBILE, AMOUNT),
    )
    md_out = os.path.join(tmp, "备忘_脱敏.md")
    code, out, _err = cli(
        [
            "redact", md_source,
            "-o", md_out,
            "--map", os.path.join(tmp, "备忘.map.json"),
            "--no-audit",
            "--no-candidates",
        ]
    )
    md_text = read_text(md_out)
    md_hits = [line for line in out.splitlines() if line.startswith("hits:")]
    checks.append(("md 脱敏返回 0", code == 0))
    checks.append(("md 脱敏件无原文", COMPANY not in md_text and MOBILE not in md_text))
    checks.append(("md 也有 hits 摘要", len(md_hits) == 1 and "公司名 1" in md_hits[0]))
    checks.append(("md hits 有电话", bool(md_hits) and "电话 1" in md_hits[0]))
    checks.append(("md 金额未动", AMOUNT in md_text))

    # ---- 3. restore 子命令 ----
    restored = os.path.join(tmp, "合同_还原.docx")
    code, _out, _err = cli(["restore", redacted, "--map", map_path, "-o", restored])
    checks.append(("restore 子命令返回 0", code == 0))
    checks.append(("还原后与原文逐字一致", docx_text(restored) == original))

    # ---- 4. 原有参数透传：--only / --mask-only ----
    only_out = os.path.join(tmp, "only.docx")
    cli(
        [
            "redact", source, "-o", only_out,
            "--map", os.path.join(tmp, "only.map.json"),
            "--only", "phone",
            "--no-audit", "--no-candidates",
        ]
    )
    only_text = docx_text(only_out)
    checks.append(("--only phone 遮了手机", MOBILE not in only_text))
    checks.append(("--only phone 保留公司名", COMPANY in only_text))

    mask_out = os.path.join(tmp, "mask.docx")
    mask_map = os.path.join(tmp, "mask.map.json")
    mask_audit = os.path.join(tmp, "mask.audit.jsonl")
    code, _out, _err = cli(
        [
            "redact", source, "-o", mask_out,
            "--map", mask_map,
            "--mask-only",
            "--audit-json", mask_audit,
            "--no-audit", "--no-candidates",
        ]
    )
    mask_text = docx_text(mask_out)
    checks.append(("--mask-only 返回 0", code == 0))
    checks.append(("掩码后没有原文", not any(secret in mask_text for secret in SECRETS)))
    checks.append(("掩码后打了星号", "*" in mask_text))
    checks.append(("掩码模式不生成映射表", not os.path.exists(mask_map)))
    mask_entry = json.loads(open(mask_audit, encoding="utf-8").read().strip().splitlines()[-1])
    checks.append(("掩码模式审计记为 mask-only", mask_entry.get("mode") == "mask-only"))
    checks.append(("掩码模式审计无映射指纹", mask_entry.get("mapping_sha256") is None))

    # ---- 5. 错误处理 ----
    code, _out, _err = cli(["redact", source, "--bogus-flag"])
    checks.append(("redact 参数错误返回 2", code == 2))
    code, _out, err = cli(["redact", source, "--audit-json"])
    checks.append(("--audit-json 缺值返回 2", code == 2))
    checks.append(("--audit-json 缺值有提示", "--audit-json" in err))

    # check 子命令冒烟：干净件返回 0
    code, out, _err = cli(["check", redacted])
    checks.append(("check 子命令可用于干净件", code == 0))
    checks.append(("check 输出干净结论", "结果：干净" in out))

    for name, ok in checks:
        print("%s %s" % ("✓" if ok else "✗", name))
    failed = [name for name, ok in checks if not ok]
    print("\n结果：%s" % ("全部通过" if not failed else "失败：" + "、".join(failed)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
