#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""验收：规则开关、不可逆掩码、加密映射表、恢复报告、审计日志。"""
import json
import os
import subprocess
import sys
import tempfile

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PY = sys.executable
SECRETS = ["星海精密制造有限公司", "陈志远", "13800000001", "Abc@12345", "sk-live-9f8e7d6c5b4a"]


def text_of(path):
    return "\n".join(p.text for p in Document(path).paragraphs)


def run(args, env=None):
    return subprocess.run([PY, os.path.join(ROOT, "redact.py")] + args,
                          capture_output=True, text=True, cwd=ROOT, env=env)


def main():
    tmp = tempfile.mkdtemp(prefix="modes-test-")
    src = os.path.join(tmp, "in.docx")
    doc = Document()
    for line in ["甲方：星海精密制造有限公司", "联系人：陈志远 手机：13800000001",
                 "密码：Abc@12345", "api_key = sk-live-9f8e7d6c5b4a",
                 "合同金额 1,234,567.89 元，违约金 20%。"]:
        doc.add_paragraph(line)
    doc.save(src)

    fails = []

    # 1. --only phone：只遮手机号
    out1 = os.path.join(tmp, "only.docx")
    run([src, "-o", out1, "--map", os.path.join(tmp, "only.json"), "--only", "phone", "--no-audit"])
    t1 = text_of(out1)
    if "13800000001" in t1:
        fails.append("--only phone 没遮手机号")
    if "星海精密制造有限公司" not in t1:
        fails.append("--only phone 误遮了公司名")

    # 2. --skip password：密码不遮
    out2 = os.path.join(tmp, "skip.docx")
    run([src, "-o", out2, "--map", os.path.join(tmp, "skip.json"), "--skip", "password", "--no-audit"])
    t2 = text_of(out2)
    if "Abc@12345" not in t2:
        fails.append("--skip password 仍然遮了密码")

    # 3. --mask-only：不可逆，不留映射表
    out3 = os.path.join(tmp, "mask.docx")
    map3 = os.path.join(tmp, "mask.json")
    run([src, "-o", out3, "--map", map3, "--mask-only", "--no-audit"])
    t3 = text_of(out3)
    if any(s in t3 for s in SECRETS):
        fails.append("掩码模式还有原文残留")
    if os.path.exists(map3):
        fails.append("掩码模式不该生成映射表")
    if "*" not in t3:
        fails.append("掩码模式没有打星")

    # 4. 口令加密 + 恢复报告
    out4 = os.path.join(tmp, "enc.docx")
    map4 = os.path.join(tmp, "enc.json")
    run([src, "-o", out4, "--map", map4, "--passphrase", "pw-123", "--no-audit"])
    blob = json.load(open(map4, encoding="utf-8"))
    if blob.get("cipher") != "hmac-ctr-sha256":
        fails.append("映射表没有加密")
    bad = subprocess.run([PY, os.path.join(ROOT, "restore.py"), out4, "--map", map4,
                          "--passphrase", "wrong", "-o", os.path.join(tmp, "bad.docx")],
                         capture_output=True, text=True)
    if bad.returncode == 0:
        fails.append("错口令竟然恢复成功")
    rep = os.path.join(tmp, "restore.json")
    ok = subprocess.run([PY, os.path.join(ROOT, "restore.py"), out4, "--map", map4,
                         "--passphrase", "pw-123", "-o", os.path.join(tmp, "back.docx"),
                         "--report", rep], capture_output=True, text=True)
    if ok.returncode != 0:
        fails.append("对的口令恢复失败")
    else:
        data = json.load(open(rep, encoding="utf-8"))
        if data["restored"] == 0 or data["unused_tokens"]:
            fails.append("恢复报告数字不对：%s" % data)

    # 5. 审计日志
    home = os.path.join(tmp, "audit-home")
    env = dict(os.environ, CONTRACT_REDACT_HOME=home)
    run([src, "-o", os.path.join(tmp, "audit.docx"), "--map", os.path.join(tmp, "audit.json")], env=env)
    log = os.path.join(home, "audit.log")
    if not os.path.isfile(log):
        fails.append("没有写审计日志")
    else:
        entry = json.loads(open(log, encoding="utf-8").read().strip().splitlines()[-1])
        if entry["entities"].get("phone") != 1:
            fails.append("审计日志统计不对：%s" % entry["entities"])

    print("规则开关/掩码/加密/报告/审计：%s" % ("全部通过" if not fails else "失败 - " + "；".join(fails)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
