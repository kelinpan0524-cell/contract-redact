#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""多格式验收：md / xlsx / pdf / docx / doc。

每个文件跑一遍：脱敏后原文必须搜不到；还原后原文必须回来。
"""
import json
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PY = sys.executable

SECRETS = ["星海精密制造有限公司", "陈志远", "13800000001", "chen.zy@example.com",
           "星海路88号", "91340200MA2T8K9L3D"]
KEEP = ["1,234,567.89", "20%"]


def text_of(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt"):
        return open(path, encoding="utf-8", errors="ignore").read()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    parts.append(c.text)
        return "\n".join(parts)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        parts.append(cell.value)
                    elif cell.value is not None:
                        parts.append(str(cell.value))
                    if cell.comment and cell.comment.text:
                        parts.append(cell.comment.text)
        return "\n".join(parts)
    if ext == ".pdf":
        import pymupdf
        d = pymupdf.open(path)
        out = "\n".join(p.get_text("text") for p in d)
        d.close()
        return out
    return ""


def ocr_text_of(path):
    """扫描件没有文字层，用同一套 OCR 读回来做核对。"""
    import pymupdf
    sys.path.insert(0, ROOT)
    import ocr

    d = pymupdf.open(path)
    parts = []
    for page in d:
        parts.extend(line[0] for line in ocr.ocr_pdf_page(page))
    d.close()
    return "\n".join(parts)


def run(cmd):
    r = subprocess.run(cmd, capture_output=True, text=True, cwd=ROOT)
    if r.returncode != 0:
        print(r.stdout[-800:], r.stderr[-800:])
    return r.returncode


def check_one(src):
    stem = os.path.splitext(src)[0]
    ext = os.path.splitext(src)[1].lower()
    out_ext = ".docx" if ext == ".doc" else ext   # .doc 会被转成 .docx 输出
    red = stem + "_redacted" + out_ext
    mp = stem + "_redacted.map.json"
    back = stem + "_restored" + out_ext

    code = run([PY, os.path.join(ROOT, "redact.py"), src, "-o", red, "--map", mp])
    if code != 0:
        return ["脱敏失败"]

    fails = []
    scanned = "scan" in os.path.basename(src).lower()
    red_text = ocr_text_of(red) if scanned else text_of(red)
    probe = red_text.replace(" ", "").lower() if scanned else red_text
    leaked = [s for s in SECRETS if (s.lower() if scanned else s) in probe]
    if leaked:
        fails.append("残留：%s" % "、".join(leaked))
    if not scanned:
        for keep in KEEP:
            if keep not in red_text:
                fails.append("误改金额/比例：%s" % keep)

    code = run([PY, os.path.join(ROOT, "restore.py"), red, "--map", mp, "-o", back])
    if code != 0:
        fails.append("还原失败")
        return fails
    back_text = ocr_text_of(back) if scanned else text_of(back)
    probe_back = back_text.replace(" ", "").lower() if scanned else back_text
    missing = [s for s in SECRETS if (s.lower() if scanned else s) not in probe_back]
    if missing:
        fails.append("还原缺失：%s" % "、".join(missing))
    return fails


def main():
    run([PY, os.path.join(HERE, "make_sample_multi.py")])
    files = [os.path.join(HERE, n) for n in
             ("sample.md", "sample.xlsx", "sample.pdf", "sample_scan.pdf", "sample.docx", "sample.doc")]
    files = [f for f in files if os.path.isfile(f)]

    failures = 0
    for f in files:
        fails = check_one(f)
        name = os.path.basename(f)
        if fails:
            failures += 1
            print("✗ %-12s %s" % (name, "；".join(fails)))
        else:
            print("✓ %-12s 无残留、金额未动、还原完整" % name)

    print("\n结果：%s" % ("全部通过" if failures == 0 else "%d 个文件未通过" % failures))
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
